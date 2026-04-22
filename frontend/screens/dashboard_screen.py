from datetime import datetime, timedelta
import csv

from PySide6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QFrame,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QProgressBar,
    QSizePolicy,
    QFileDialog,
)
from PySide6.QtCore import Qt, QThread, Signal, QSize

from qfluentwidgets import (
    SmoothScrollArea,
    ElevatedCardWidget,
    BodyLabel,
    CaptionLabel,
    PushButton,
    PrimaryPushButton,
    TransparentPushButton,
    FluentIcon as FIF,
)

from components.cards import StatCard, SectionHeader
from components.donut import DonutChart
from components.loading import LoadingOverlay
from components.toast import toast_error, toast_success
from ui.icon_manager import IconManager
from ui.theme import CARD_ACCENT
import api.stats_api as stats_api
import api.audit_api as audit_api
import api.schedule_api as schedule_api


# ── Worker ───────────────────────────────────────────────────────
class DashboardWorker(QThread):
    finished = Signal(dict, list, list, list, list, dict, list, str)
    error = Signal(str)

    def run(self):
        try:
            overview = stats_api.get_overview()
            by_dept = stats_api.get_by_department()
            by_degree   = stats_api.get_by_degree()
            by_position = stats_api.get_by_position()
            lecturer_status = stats_api.get_lecturer_status()
            logs = audit_api.get_audit_logs(size=8)
            recent_logs = logs.get("items", []) if isinstance(logs, dict) else []

            today_key = datetime.now().strftime("%a")
            today_payload = schedule_api.get_schedules(page=1, size=16, day_of_week=today_key)
            today_schedules = today_payload.get("items", []) if isinstance(today_payload, dict) else []
            today_schedules = sorted(today_schedules, key=lambda x: x.get("start_time") or "")

            updated_at = datetime.now().strftime("%H:%M")
            self.finished.emit(
                overview,
                by_dept,
                by_degree,
                by_position,
                recent_logs,
                lecturer_status,
                today_schedules,
                updated_at,
            )
        except Exception as e:
            self.error.emit(str(e))


# ── Screen ────────────────────────────────────────────────────────
class DashboardScreen(SmoothScrollArea):
    """Dashboard — stat cards, bar charts, recent activity."""

    def __init__(self, user_info: dict, parent=None):
        super().__init__(parent=parent)
        self.setObjectName("DashboardScreen")
        self.user_info = user_info
        self.is_admin = user_info.get("role") == "admin"
        self._worker = None
        self._latest_data: dict = {}

        self._content = QWidget()
        self._content.setObjectName("dashboardContent")
        self.setWidget(self._content)
        self.setWidgetResizable(True)

        self._build_ui()
        self._loading = LoadingOverlay(self)

    def _build_ui(self):
        layout = QVBoxLayout(self._content)
        layout.setContentsMargins(28, 24, 28, 28)
        layout.setSpacing(20)

        header_row = QHBoxLayout()
        header_row.setSpacing(12)
        header = SectionHeader(
            "Dashboard",
            "Tổng quan hệ thống — cập nhật theo thời gian thực",
            icon="squares-four",
        )
        header_row.addWidget(header, stretch=1)

        header_actions = QVBoxLayout()
        header_actions.setSpacing(6)
        header_actions.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self._updated_lbl = CaptionLabel("Cập nhật lúc --:--")
        self._updated_lbl.setStyleSheet("color:#8B949E;")

        self._refresh_btn = PushButton(FIF.SYNC, "  Làm mới", self)
        self._refresh_btn.clicked.connect(self.refresh)

        header_actions.addWidget(self._updated_lbl, alignment=Qt.AlignmentFlag.AlignRight)
        header_actions.addWidget(self._refresh_btn, alignment=Qt.AlignmentFlag.AlignRight)
        header_row.addLayout(header_actions)
        layout.addLayout(header_row)

        # ── Stat cards row ────────────────────────────────────────
        cards_row = QHBoxLayout()
        cards_row.setSpacing(16)

        acc = CARD_ACCENT
        self._card_lect  = StatCard("users",          "—", "Tổng Giảng Viên",  acc["lecturers"])
        self._card_dept  = StatCard("buildings",       "—", "Số Khoa / Bộ Môn", acc["departments"])
        self._card_sched = StatCard("calendar-blank",  "—", "Lịch Giảng Dạy",   acc["schedules"])
        self._card_acc   = StatCard("shield-check",    "—", "Tài Khoản",         acc["accounts"])

        self._card_lect.clicked.connect(lambda: self._navigate_to("screen_lecturers"))
        self._card_dept.clicked.connect(lambda: self._navigate_to("screen_departments"))
        self._card_sched.clicked.connect(lambda: self._navigate_to("screen_schedules"))
        self._card_acc.clicked.connect(lambda: self._navigate_to("screen_accounts"))
        if not self.is_admin:
            self._card_acc.setCursor(Qt.CursorShape.ArrowCursor)

        for card in [self._card_lect, self._card_dept, self._card_sched, self._card_acc]:
            cards_row.addWidget(card)
        layout.addLayout(cards_row)

        # ── Charts row ────────────────────────────────────────────
        charts_row = QHBoxLayout()
        charts_row.setSpacing(16)

        # Left: department table
        self._dept_card = self._make_card(
            "buildings", "Giảng Viên theo Khoa", stretch=3
        )
        dept_inner_v = self._dept_card.property("inner_layout")
        self._dept_table = QTableWidget(0, 3)
        self._dept_table.setHorizontalHeaderLabels(["Tên Khoa", "Mã", "Số GV"])
        self._dept_table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeMode.Stretch
        )
        self._dept_table.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.ResizeMode.Fixed
        )
        self._dept_table.horizontalHeader().setSectionResizeMode(
            2, QHeaderView.ResizeMode.Fixed
        )
        self._dept_table.setColumnWidth(1, 70)
        self._dept_table.setColumnWidth(2, 70)
        self._dept_table.verticalHeader().setVisible(False)
        self._dept_table.setSelectionMode(QTableWidget.SelectionMode.NoSelection)
        self._dept_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self._dept_table.setAlternatingRowColors(True)
        self._dept_table.setMinimumHeight(160)
        self._dept_table.setMaximumHeight(340)
        self._dept_table.setShowGrid(False)
        dept_inner_v.addWidget(self._dept_table)
        charts_row.addWidget(self._dept_card, stretch=3)

        # Right: degree bar chart
        self._degree_card = self._make_card(
            "graduation-cap", "Giảng Viên theo Trình Độ", stretch=2
        )
        self._degree_inner = self._degree_card.property("inner_layout")
        charts_row.addWidget(self._degree_card, stretch=2)

        layout.addLayout(charts_row)

        # ── Second row ────────────────────────────────────────────
        second_row = QHBoxLayout()
        second_row.setSpacing(16)

        # Position chart
        self._pos_card = self._make_card("users", "Giảng Viên theo Chức Vụ", stretch=2)
        self._pos_inner = self._pos_card.property("inner_layout")
        second_row.addWidget(self._pos_card, stretch=2)

        # Recent activity (admin only)
        if self.is_admin:
            self._act_card = self._make_card("clipboard-text", "Hoạt Động Gần Đây", stretch=3)
            self._act_inner = self._act_card.property("inner_layout")

            act_scroll = SmoothScrollArea()
            act_scroll.setWidgetResizable(True)
            act_scroll.setFrameShape(QFrame.Shape.NoFrame)
            act_scroll.setStyleSheet(
                "SmoothScrollArea{background:transparent;border:none;}"
                "QScrollArea{background:transparent;border:none;}"
            )
            act_scroll.setMinimumHeight(200)

            self._act_content = QWidget()
            self._act_content.setStyleSheet("background:transparent;")
            self._act_vbox = QVBoxLayout(self._act_content)
            self._act_vbox.setContentsMargins(0, 0, 0, 0)
            self._act_vbox.setSpacing(0)
            self._act_vbox.addStretch()
            act_scroll.setWidget(self._act_content)
            act_scroll.viewport().setStyleSheet("background:transparent;")

            self._act_inner.addWidget(act_scroll, stretch=1)

            more_btn = TransparentPushButton(FIF.HISTORY, "  Nhật ký đầy đủ →", self._act_card)
            more_btn.clicked.connect(lambda: self._navigate_to("screen_audit"))
            self._act_inner.addWidget(more_btn, alignment=Qt.AlignmentFlag.AlignRight)
            second_row.addWidget(self._act_card, stretch=3)

        layout.addLayout(second_row)

        third_row = QHBoxLayout()
        third_row.setSpacing(16)

        self._today_card = self._make_card("calendar-check", "Lịch Hôm Nay", stretch=2)
        self._today_inner = self._today_card.property("inner_layout")
        self._today_hint = CaptionLabel("Danh sách ca dạy theo thứ hiện tại")
        self._today_hint.setStyleSheet("color:#8B949E;")
        self._today_inner.addWidget(self._today_hint)

        today_scroll = SmoothScrollArea()
        today_scroll.setWidgetResizable(True)
        today_scroll.setFrameShape(QFrame.Shape.NoFrame)
        today_scroll.setStyleSheet(
            "SmoothScrollArea{background:transparent;border:none;}"
            "QScrollArea{background:transparent;border:none;}"
        )
        today_scroll.setMinimumHeight(160)
        self._today_content = QWidget()
        self._today_content.setStyleSheet("background:transparent;")
        self._today_vbox = QVBoxLayout(self._today_content)
        self._today_vbox.setContentsMargins(0, 0, 0, 0)
        self._today_vbox.setSpacing(0)
        self._today_vbox.addStretch()
        today_scroll.setWidget(self._today_content)
        today_scroll.viewport().setStyleSheet("background:transparent;")
        self._today_inner.addWidget(today_scroll, stretch=1)
        third_row.addWidget(self._today_card, stretch=2)

        # Lecturer status donut card
        self._status_card = self._make_card("users", "Trạng Thái Giảng Viên")
        _si = self._status_card.property("inner_layout")

        donut_row = QHBoxLayout()
        donut_row.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._status_donut = DonutChart(size=100, hole=0.55)
        donut_row.addWidget(self._status_donut)
        _si.addLayout(donut_row)

        _STATUS_ITEMS = [
            ("Đang dạy",  "#3FB950"),
            ("Tạm nghỉ",  "#D29922"),
            ("Nghỉ việc", "#F85149"),
        ]
        self._status_lbls: dict = {}
        for label, color in _STATUS_ITEMS:
            leg_h = QHBoxLayout()
            leg_h.setSpacing(8)
            dot = QLabel("●")
            dot.setStyleSheet(f"color:{color}; font-size:10px; background:transparent;")
            lbl = CaptionLabel(f"{label}: —")
            lbl.setStyleSheet("background:transparent;")
            leg_h.addWidget(dot)
            leg_h.addWidget(lbl, stretch=1)
            self._status_lbls[label] = lbl
            _si.addLayout(leg_h)
        _si.addStretch()
        third_row.addWidget(self._status_card, stretch=2)

        if self.is_admin:
            self._quick_card = self._make_card("rocket-launch", "Thao Tác Nhanh", stretch=2)
            self._quick_inner = self._quick_card.property("inner_layout")
            quick_btn_row = QHBoxLayout()
            quick_btn_row.setSpacing(8)

            add_lect_btn = PrimaryPushButton(FIF.ADD, "  Thêm Giảng Viên", self)
            add_lect_btn.clicked.connect(
                lambda: self._navigate_to(
                    "screen_lecturers",
                    after=lambda w: getattr(w, "_on_add", lambda: None)(),
                )
            )
            add_sched_btn = PrimaryPushButton(FIF.CALENDAR, "  Thêm Lịch Dạy", self)
            add_sched_btn.clicked.connect(
                lambda: self._navigate_to(
                    "screen_schedules",
                    after=lambda w: getattr(w, "_on_add", lambda: None)(),
                )
            )
            export_btn = PushButton(FIF.DOWNLOAD, "  Xuất Báo Cáo", self)
            export_btn.clicked.connect(self._export_dashboard_report)

            quick_btn_row.addWidget(add_lect_btn)
            quick_btn_row.addWidget(add_sched_btn)
            quick_btn_row.addWidget(export_btn)
            self._quick_inner.addLayout(quick_btn_row)
            third_row.addWidget(self._quick_card, stretch=2)

        layout.addLayout(third_row)
        layout.addStretch()

    # ── Card factory ─────────────────────────────────────────────
    def _make_card(self, icon_name: str, title: str, stretch: int = 1) -> ElevatedCardWidget:
        card = ElevatedCardWidget()
        card.setMinimumHeight(230)
        card.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding
        )
        card_v = QVBoxLayout(card)
        card_v.setContentsMargins(16, 16, 16, 16)
        card_v.setSpacing(10)

        title_row = QHBoxLayout()
        icon_lbl = QLabel()
        icon_lbl.setStyleSheet("background:transparent;")
        icon_lbl.setPixmap(IconManager.get(icon_name, "#8B949E", 18).pixmap(18, 18))
        t_lbl = BodyLabel(title)
        t_lbl.setStyleSheet("font-size:14px; font-weight:600; background:transparent;")
        title_row.addWidget(icon_lbl)
        title_row.addSpacing(6)
        title_row.addWidget(t_lbl)
        title_row.addStretch()
        card_v.addLayout(title_row)

        inner = QVBoxLayout()
        inner.setSpacing(8)
        card_v.addLayout(inner, stretch=1)

        card.setProperty("inner_layout", inner)
        return card

    # ── Refresh ────────────────────────────────────────────────────
    def refresh(self):
        if self._worker is not None and self._worker.isRunning():
            return
        self._loading.show()
        self._worker = DashboardWorker()
        self._worker.finished.connect(self._on_loaded)
        self._worker.error.connect(self._on_error)
        self._worker.start()

    def _on_loaded(
        self,
        overview: dict,
        by_dept: list,
        by_degree: list,
        by_position: list,
        recent_logs: list,
        lecturer_status: dict,
        today_schedules: list,
        updated_at: str,
    ):
        self._loading.hide()
        self._latest_data = {
            "overview": overview,
            "by_dept": by_dept,
            "by_degree": by_degree,
            "by_position": by_position,
            "lecturer_status": lecturer_status,
            "today_schedules": today_schedules,
            "recent_logs": recent_logs,
            "updated_at": updated_at,
        }
        self._updated_lbl.setText(f"Cập nhật lúc {updated_at}")

        # Stat cards
        self._card_lect.set_value(str(overview.get("total_lecturers", 0)))
        self._card_dept.set_value(str(overview.get("total_departments", 0)))
        self._card_sched.set_value(str(overview.get("total_schedules", 0)))
        self._card_acc.set_value(str(overview.get("total_accounts", 0)))
        active   = lecturer_status.get("active", 0)
        on_leave = lecturer_status.get("on_leave", 0)
        resigned = lecturer_status.get("resigned", 0)
        self._card_lect.set_meta_text(
            f"Đang dạy: {active}   •   Tạm nghỉ: {on_leave}   •   Nghỉ việc: {resigned}",
            color="#6E7681",
        )
        # Donut + legend
        self._status_donut.set_data([
            ("Đang dạy",  active,   "#3FB950"),
            ("Tạm nghỉ",  on_leave, "#D29922"),
            ("Nghỉ việc", resigned, "#F85149"),
        ])
        for label, count in [("Đang dạy", active), ("Tạm nghỉ", on_leave), ("Nghỉ việc", resigned)]:
            if label in self._status_lbls:
                self._status_lbls[label].setText(f"{label}: {count}")

        # Dept table
        self._dept_table.setRowCount(0)
        for row_data in by_dept:
            r = self._dept_table.rowCount()
            self._dept_table.insertRow(r)
            self._dept_table.setItem(r, 0, QTableWidgetItem(row_data.get("department_name", "")))
            self._dept_table.setItem(r, 1, QTableWidgetItem(row_data.get("department_code", "")))
            cnt = QTableWidgetItem(str(row_data.get("count", 0)))
            cnt.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self._dept_table.setItem(r, 2, cnt)

        # Bar charts
        self._render_bar_chart(self._degree_inner, by_degree, "degree", "#0099FF")
        self._render_bar_chart(self._pos_inner,    by_position, "position", "#8764B8")

        # Activity feed
        if self.is_admin and hasattr(self, "_act_vbox"):
            # Xóa các row cũ, giữ lại stretch ở cuối
            while self._act_vbox.count() > 1:
                it = self._act_vbox.takeAt(0)
                if it.widget():
                    it.widget().hide()
                    it.widget().deleteLater()

            ACTION_COLORS = {
                "login": "#58A6FF", "create": "#3FB950",
                "update": "#D29922", "delete": "#F85149",
            }
            ACTION_VI = {
                "login": "Đăng nhập", "create": "Tạo mới",
                "update": "Cập nhật",  "delete": "Xóa",
            }
            ENTITY_VI = {
                "lecturer": "giảng viên", "department": "khoa",
                "schedule": "lịch dạy",  "account": "tài khoản",
            }
            ENTITY_COLOR = {
                "lecturer": "#3FB950", "schedule": "#8764B8",
                "account": "#D29922",  "department": "#58A6FF",
            }

            for i, log in enumerate(recent_logs):
                action     = log.get("action", "")
                color      = ACTION_COLORS.get(action, "#8B949E")
                user       = log.get("username", "—")
                entity     = log.get("entity_type") or ""
                rel_t      = self._to_relative_time(log.get("created_at", ""))
                chip_color = ENTITY_COLOR.get(entity, "#6E7681")

                row_w = QFrame()
                row_w.setFixedHeight(44)
                row_w.setStyleSheet(
                    "QFrame{background:transparent;border:none;"
                    "border-bottom:1px solid #21262D;}"
                )
                row_h = QHBoxLayout(row_w)
                row_h.setContentsMargins(4, 0, 4, 0)
                row_h.setSpacing(10)

                # Avatar chữ tắt
                initials = user[:2].upper() if user and user != "—" else "?"
                avatar = QLabel(initials)
                avatar.setFixedSize(30, 30)
                avatar.setAlignment(Qt.AlignmentFlag.AlignCenter)
                ar, ag, ab = int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16)
                avatar.setStyleSheet(
                    f"background:rgba({ar},{ag},{ab},55); border-radius:15px;"
                    f"color:{color}; font-size:11px; font-weight:700;"
                )

                # Chip entity
                chip = QLabel(ENTITY_VI.get(entity, entity).title())
                chip.setFixedHeight(20)
                chip.setStyleSheet(
                    f"color:white; background:{chip_color}; border-radius:9px;"
                    f"padding:0px 8px; font-size:11px;"
                )

                # Nội dung
                text_l = QLabel(f"{user}  •  {ACTION_VI.get(action, action)}")
                text_l.setStyleSheet(f"color:{color}; font-size:12px;")

                # Thời gian
                time_l = QLabel(rel_t)
                time_l.setFixedWidth(90)
                time_l.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                time_l.setStyleSheet("color:#6E7681; font-size:11px;")

                row_h.addWidget(avatar)
                row_h.addWidget(chip)
                row_h.addWidget(text_l, stretch=1)
                row_h.addWidget(time_l)

                self._act_vbox.insertWidget(i, row_w)

        self._render_today_schedules(today_schedules)

    def _render_bar_chart(self, layout: QVBoxLayout, data: list, key: str, color: str):
        while layout.count():
            it = layout.takeAt(0)
            if it.widget():
                it.widget().hide()
                it.widget().deleteLater()

        if not data:
            lbl = CaptionLabel("Không có dữ liệu")
            lbl.setStyleSheet("color:#484F58;")
            layout.addWidget(lbl)
            return

        total = sum(d.get("count", 0) for d in data) or 1

        for d in data[:8]:
            name  = str(d.get(key, "—"))
            count = d.get("count", 0)
            pct   = int(count / total * 100)

            row_w  = QWidget()
            row_w.setStyleSheet("background:transparent;")
            row_h = QHBoxLayout(row_w)
            row_h.setContentsMargins(0, 0, 0, 0)
            row_h.setSpacing(8)

            name_lbl = QLabel(name)
            name_lbl.setFixedWidth(130)
            name_lbl.setStyleSheet("color:#8B949E; font-size:12px; background:transparent;")
            name_lbl.setWordWrap(False)

            bar = QProgressBar()
            bar.setRange(0, 100)
            bar.setValue(pct)
            bar.setFixedHeight(8)
            bar.setTextVisible(False)
            bar.setStyleSheet(
                f"QProgressBar{{background:#21262D;border:none;border-radius:4px;}}"
                f"QProgressBar::chunk{{background:{color};border-radius:4px;}}"
            )

            pct_lbl = QLabel(f"{pct}%")
            pct_lbl.setFixedWidth(30)
            pct_lbl.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            pct_lbl.setStyleSheet("color:#484F58; font-size:11px; background:transparent;")

            cnt_lbl = QLabel(str(count))
            cnt_lbl.setFixedWidth(30)
            cnt_lbl.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            cnt_lbl.setStyleSheet(f"color:{color}; font-size:12px; font-weight:600; background:transparent;")

            row_h.addWidget(name_lbl)
            row_h.addWidget(bar, stretch=1)
            row_h.addWidget(pct_lbl)
            row_h.addWidget(cnt_lbl)
            layout.addWidget(row_w)

    def _on_error(self, msg: str):
        self._loading.hide()
        for card in [self._card_lect, self._card_dept, self._card_sched, self._card_acc]:
            card.set_value("!")
        toast_error(self.window(), msg)

    def _navigate_to(self, widget_attr: str, after=None):
        if widget_attr == "screen_accounts" and not self.is_admin:
            return

        win = self.window()
        target = getattr(win, widget_attr, None)
        if not target:
            return

        if hasattr(win, "switchTo"):
            win.switchTo(target)

        if callable(after):
            after(target)

    def _to_relative_time(self, iso_text: str) -> str:
        if not iso_text:
            return "vừa xong"

        try:
            clean = str(iso_text).replace("Z", "")
            dt = datetime.fromisoformat(clean)
        except Exception:
            return str(iso_text)[:16].replace("T", " ")

        now = datetime.now()
        delta = now - dt

        if delta < timedelta(minutes=1):
            return "vừa xong"
        if delta < timedelta(hours=1):
            return f"{max(1, int(delta.total_seconds() // 60))} phút trước"
        if delta < timedelta(hours=24) and dt.date() == now.date():
            return f"{max(1, int(delta.total_seconds() // 3600))} giờ trước"

        if dt.date() == (now.date() - timedelta(days=1)):
            return f"hôm qua lúc {dt.strftime('%H:%M')}"

        return dt.strftime("%d/%m %H:%M")

    def _render_today_schedules(self, today_schedules: list):
        while self._today_vbox.count() > 1:
            it = self._today_vbox.takeAt(0)
            if it.widget():
                it.widget().hide()
                it.widget().deleteLater()

        if not today_schedules:
            empty_w = QFrame()
            empty_w.setFixedHeight(44)
            empty_w.setStyleSheet("QFrame{background:transparent;border:none;}")
            empty_h = QHBoxLayout(empty_w)
            empty_h.setContentsMargins(4, 0, 4, 0)
            empty_lbl = CaptionLabel("Không có ca dạy trong hôm nay")
            empty_lbl.setStyleSheet("color:#484F58;")
            empty_h.addWidget(empty_lbl)
            self._today_vbox.insertWidget(0, empty_w)
            return

        for i, row in enumerate(today_schedules[:8]):
            lecturer_name = ((row.get("lecturer") or {}).get("full_name") or "—").strip()
            start_t = row.get("start_time", "--:--")
            end_t   = row.get("end_time",   "--:--")
            subject = row.get("subject_name", "—")
            code    = row.get("subject_code", "")
            room    = row.get("room", "—")

            row_w = QFrame()
            row_w.setFixedHeight(60)
            row_w.setStyleSheet(
                "QFrame{background:transparent;border:none;"
                "border-bottom:1px solid #21262D;}"
            )
            row_h = QHBoxLayout(row_w)
            row_h.setContentsMargins(4, 4, 4, 4)
            row_h.setSpacing(10)

            # Time block
            time_w = QWidget()
            time_w.setFixedWidth(80)
            time_w.setStyleSheet(
                "background:rgba(88,166,255,30);border-radius:6px;"
            )
            time_v = QVBoxLayout(time_w)
            time_v.setContentsMargins(4, 2, 4, 2)
            time_v.setSpacing(0)
            time_v.setAlignment(Qt.AlignmentFlag.AlignCenter)
            t1 = QLabel(start_t)
            t1.setAlignment(Qt.AlignmentFlag.AlignCenter)
            t1.setStyleSheet(
                "font-size:12px;font-weight:700;color:#58A6FF;background:transparent;"
            )
            t2 = QLabel(end_t)
            t2.setAlignment(Qt.AlignmentFlag.AlignCenter)
            t2.setStyleSheet("font-size:10px;color:#8B949E;background:transparent;")
            time_v.addWidget(t1)
            time_v.addWidget(t2)

            # Subject + lecturer info
            info_v = QVBoxLayout()
            info_v.setSpacing(2)
            info_v.setContentsMargins(0, 0, 0, 0)
            subj_lbl = QLabel(f"{subject} ({code})" if code else subject)
            subj_lbl.setStyleSheet(
                "font-size:12px;font-weight:600;color:#C9D1D9;background:transparent;"
            )
            gv_lbl = QLabel(f"GV: {lecturer_name}   •   Phòng: {room}")
            gv_lbl.setStyleSheet("font-size:11px;color:#8B949E;background:transparent;")
            info_v.addWidget(subj_lbl)
            info_v.addWidget(gv_lbl)

            row_h.addWidget(time_w)
            row_h.addLayout(info_v, stretch=1)

            self._today_vbox.insertWidget(i, row_w)

    def _export_dashboard_report(self):
        if not self._latest_data:
            toast_error(self.window(), "Chưa có dữ liệu để xuất báo cáo")
            return

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Lưu báo cáo Dashboard",
            f"dashboard_report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            "Excel (*.xlsx);;CSV (*.csv)",
        )
        if not path:
            return

        try:
            if path.lower().endswith(".xlsx"):
                self._export_xlsx(path)
            else:
                self._export_csv(path)
            toast_success(self.window(), "Xuất báo cáo thành công")
        except Exception as e:
            toast_error(self.window(), str(e))

    def _export_xlsx(self, path: str):
        import openpyxl
        from openpyxl.styles import Font, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tổng Quan"

        ws.merge_cells("A1:B1")
        ws["A1"] = f"Báo cáo Dashboard — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws["A1"].font = Font(bold=True, size=13)
        ws.row_dimensions[1].height = 24

        overview = self._latest_data.get("overview", {})
        status   = self._latest_data.get("lecturer_status", {})

        ws.append([])
        header_row = ws.append(["Chỉ số", "Giá trị"])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for label, val in [
            ("Tổng giảng viên",   overview.get("total_lecturers", 0)),
            ("  Đang dạy",        status.get("active", 0)),
            ("  Tạm nghỉ",        status.get("on_leave", 0)),
            ("  Nghỉ việc",       status.get("resigned", 0)),
            ("Số khoa / bộ môn",  overview.get("total_departments", 0)),
            ("Lịch giảng dạy",    overview.get("total_schedules", 0)),
            ("Tài khoản",         overview.get("total_accounts", 0)),
        ]:
            ws.append([label, val])
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 12

        # Sheet 2 — by department
        ws2 = wb.create_sheet("Theo Khoa")
        ws2.append(["Tên Khoa", "Mã Khoa", "Số Giảng Viên"])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
        for d in self._latest_data.get("by_dept", []):
            ws2.append([d.get("department_name", ""), d.get("department_code", ""), d.get("count", 0)])
        ws2.column_dimensions["A"].width = 30
        ws2.column_dimensions["B"].width = 12
        ws2.column_dimensions["C"].width = 16

        # Sheet 3 — by degree
        ws3 = wb.create_sheet("Theo Trình Độ")
        ws3.append(["Trình Độ", "Số Giảng Viên"])
        for cell in ws3[1]:
            cell.font = Font(bold=True)
        for d in self._latest_data.get("by_degree", []):
            ws3.append([d.get("degree", ""), d.get("count", 0)])
        ws3.column_dimensions["A"].width = 12
        ws3.column_dimensions["B"].width = 16

        # Sheet 4 — by position
        ws4 = wb.create_sheet("Theo Chức Vụ")
        ws4.append(["Chức Vụ", "Số Giảng Viên"])
        for cell in ws4[1]:
            cell.font = Font(bold=True)
        for d in self._latest_data.get("by_position", []):
            ws4.append([d.get("position", "") or "Chưa phân công", d.get("count", 0)])
        ws4.column_dimensions["A"].width = 26
        ws4.column_dimensions["B"].width = 16

        wb.save(path)

    def _export_csv(self, path: str):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            overview = self._latest_data.get("overview", {})
            status   = self._latest_data.get("lecturer_status", {})
            writer.writerow(["Mục", "Giá trị"])
            writer.writerow(["Tổng giảng viên", overview.get("total_lecturers", 0)])
            writer.writerow(["Số khoa / bộ môn", overview.get("total_departments", 0)])
            writer.writerow(["Lịch giảng dạy", overview.get("total_schedules", 0)])
            writer.writerow(["Tài khoản", overview.get("total_accounts", 0)])
            writer.writerow([])
            writer.writerow(["Trạng thái", "Số lượng"])
            writer.writerow(["Đang dạy",  status.get("active", 0)])
            writer.writerow(["Tạm nghỉ",  status.get("on_leave", 0)])
            writer.writerow(["Nghỉ việc", status.get("resigned", 0)])

    def showEvent(self, event):
        super().showEvent(event)
        if not getattr(self, "_loaded", False):
            self._loaded = True
            self.refresh()
    def resizeEvent(self, event):
        self._loading.resize(self.size())
        super().resizeEvent(event)
