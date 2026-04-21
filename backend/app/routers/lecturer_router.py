from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.deps import get_current_user, require_admin
from app.schemas.lecturer import LecturerCreate, LecturerUpdate
from app.services.audit_service import log_action
from app.services.store import STORE

router = APIRouter(prefix="/lecturers", tags=["Lecturers"])


@router.get("")
def list_lecturers(
    page: int = 1,
    size: int = 20,
    search: str = "",
    department_id: int | None = None,
    degree: str = "",
    position: str = "",
    gender: str = "",
    status_filter: str = Query("", alias="status"),
    _: dict = Depends(get_current_user),
) -> dict:
    return STORE.list_lecturers(
        page=page,
        size=size,
        search=search,
        department_id=department_id,
        degree=degree,
        position=position,
        gender=gender,
        status=status_filter,
    )


@router.post("")
def create_lecturer(payload: LecturerCreate, admin: dict = Depends(require_admin)) -> dict:
    try:
        row = STORE.create_lecturer(payload.model_dump())
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc))
    log_action("create", "lecturer", admin["username"], f"Tạo giảng viên {row['employee_code']}", row["id"])
    return row


@router.put("/{lecturer_id}")
def update_lecturer(lecturer_id: int, payload: LecturerUpdate, admin: dict = Depends(require_admin)) -> dict:
    try:
        row = STORE.update_lecturer(lecturer_id, payload.model_dump(exclude_unset=True))
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc))
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy giảng viên.")
    log_action("update", "lecturer", admin["username"], f"Cập nhật giảng viên {lecturer_id}", lecturer_id)
    return row


@router.delete("/{lecturer_id}")
def delete_lecturer(lecturer_id: int, admin: dict = Depends(require_admin)) -> dict[str, str]:
    ok = STORE.delete_lecturer(lecturer_id)
    if not ok:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy giảng viên.")
    log_action("delete", "lecturer", admin["username"], f"Xóa giảng viên {lecturer_id}", lecturer_id)
    return {"message": "Xóa giảng viên thành công."}


@router.post("/{lecturer_id}/avatar")
def upload_avatar(lecturer_id: int, _: dict = Depends(require_admin)) -> dict[str, str]:
    # Placeholder de giu hop dong API voi frontend.
    if not STORE.get_lecturer(lecturer_id):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy giảng viên.")
    return {"message": "Upload avatar thành công."}


@router.get("/export")
def export_lecturers_excel(
    search: str = "",
    department_id: int | None = None,
    degree: str = "",
    position: str = "",
    gender: str = "",
    status_filter: str = Query("", alias="status"),
    _: dict = Depends(get_current_user),
) -> StreamingResponse:
    data = STORE.list_lecturers(
        page=1, size=10000,
        search=search, department_id=department_id,
        degree=degree, position=position,
        gender=gender, status=status_filter,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách giảng viên"

    headers = ["Mã GV", "Họ và tên", "Email", "SĐT", "Giới tính",
               "Ngày sinh", "Học vị", "Chức vụ", "Khoa/Bộ môn",
               "Ngày vào làm", "Trạng thái"]
    header_fill = PatternFill("solid", fgColor="1F6FEB")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    STATUS_MAP = {"active": "Đang dạy", "inactive": "Nghỉ việc", "on_leave": "Nghỉ phép"}
    GENDER_MAP = {"male": "Nam", "female": "Nữ", "other": "Khác"}
    for r_idx, row in enumerate(data["items"], 2):
        dep_name = (row.get("department") or {}).get("name", "")
        ws.append([
            row.get("employee_code", ""),
            row.get("full_name", ""),
            row.get("email", ""),
            row.get("phone", "") or "",
            GENDER_MAP.get(row.get("gender", ""), row.get("gender", "")),
            row.get("date_of_birth", "") or "",
            row.get("degree", ""),
            row.get("position", "") or "",
            dep_name,
            row.get("hire_date", "") or "",
            STATUS_MAP.get(row.get("status", ""), row.get("status", "")),
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"lecturers_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/pdf")
def export_lecturers_pdf(
    search: str = "",
    department_id: int | None = None,
    degree: str = "",
    position: str = "",
    gender: str = "",
    status_filter: str = Query("", alias="status"),
    _: dict = Depends(get_current_user),
) -> StreamingResponse:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os

    data = STORE.list_lecturers(
        page=1, size=10000,
        search=search, department_id=department_id,
        degree=degree, position=position,
        gender=gender, status=status_filter,
    )

    font_candidates = [
        "C:/Windows/Fonts/Arial.ttf",
        "C:/Windows/Fonts/calibri.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ]
    font_name = "Helvetica"
    for fp in font_candidates:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont("ViFont", fp))
                font_name = "ViFont"
            except Exception:
                pass
            break

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    title_style = ParagraphStyle("title", fontName=font_name, fontSize=14,
                                  spaceAfter=10, alignment=1)
    cell_style  = ParagraphStyle("cell",  fontName=font_name, fontSize=8,
                                  leading=10)

    STATUS_MAP = {"active": "Dang day", "inactive": "Nghi viec", "on_leave": "Nghi phep"}
    GENDER_MAP = {"male": "Nam", "female": "Nu", "other": "Khac"}

    header = ["Ma GV", "Ho va ten", "Email", "SDT", "Gioi tinh",
              "Hoc vi", "Chuc vu", "Khoa/Bo mon", "Trang thai"]
    rows = [header]
    for row in data["items"]:
        dep_name = (row.get("department") or {}).get("name", "")
        rows.append([
            Paragraph(row.get("employee_code", ""), cell_style),
            Paragraph(row.get("full_name", ""), cell_style),
            Paragraph(row.get("email", ""), cell_style),
            Paragraph(row.get("phone", "") or "", cell_style),
            GENDER_MAP.get(row.get("gender", ""), row.get("gender", "")),
            row.get("degree", ""),
            Paragraph(row.get("position", "") or "", cell_style),
            Paragraph(dep_name, cell_style),
            STATUS_MAP.get(row.get("status", ""), row.get("status", "")),
        ])

    col_widths = [2.2*cm, 4.5*cm, 5.5*cm, 2.8*cm, 2.0*cm, 2.2*cm, 3.0*cm, 4.0*cm, 2.2*cm]
    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND",  (0, 0), (-1, 0), colors.HexColor("#1F6FEB")),
        ("TEXTCOLOR",   (0, 0), (-1, 0), colors.white),
        ("FONTNAME",    (0, 0), (-1, 0), font_name),
        ("FONTSIZE",    (0, 0), (-1, 0), 9),
        ("FONTNAME",    (0, 1), (-1, -1), font_name),
        ("FONTSIZE",    (0, 1), (-1, -1), 8),
        ("ALIGN",       (0, 0), (-1, 0), "CENTER"),
        ("VALIGN",      (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
        ("GRID",        (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7DE")),
        ("TOPPADDING",  (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
    ]))

    elements = [
        Paragraph("DANH SACH GIANG VIEN", title_style),
        Spacer(1, 0.3*cm),
        table,
    ]
    doc.build(elements)
    buf.seek(0)
    filename = f"lecturers_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{lecturer_id}")
def get_lecturer(lecturer_id: int, _: dict = Depends(get_current_user)) -> dict:
    row = STORE.get_lecturer(lecturer_id)
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy giảng viên.")
    return row
